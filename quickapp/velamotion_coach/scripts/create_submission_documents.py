#!/usr/bin/env python3
"""Generate contest PDF/DOCX from the reviewed project content and real captures.
Dependencies: python-docx, reportlab (document generation only).
"""
from pathlib import Path
import os
import hashlib
from html import escape
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, Table, TableStyle, PageBreak
from reportlab.lib.pagesizes import A4

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'docs'
FONT = Path(os.environ.get('VELAMOTION_DOC_FONT', str(Path.home() / '.vela/sdk/qa/font/MiSansW_Regular.ttf')))
if not FONT.exists():
    raise SystemExit('Set VELAMOTION_DOC_FONT to a Chinese/Latin TrueType font (e.g. MiSansW_Regular.ttf).')
pdfmetrics.registerFont(TTFont('Chinese', str(FONT)))
styles = getSampleStyleSheet()
styles.add(ParagraphStyle('CN', fontName='Chinese', fontSize=10, leading=17, spaceAfter=9, wordWrap='CJK'))
styles.add(ParagraphStyle('CNTitle', parent=styles['CN'], fontSize=23, leading=31, textColor=colors.HexColor('#16755A'), spaceAfter=15))
styles.add(ParagraphStyle('CNHeading', parent=styles['CN'], fontSize=14, leading=22, textColor=colors.HexColor('#16755A'), spaceBefore=8))
styles.add(ParagraphStyle('CNCaption', parent=styles['CN'], fontSize=8, leading=12))
pages = [
('腕动教练 · VelaMotion Coach', [
 ('p','2026 首届 openvela AI 硬件开发者大赛｜手表应用创新方向\n队伍：DDLqudong（289）｜版本：1.0.0｜2026-09-20'),
 ('h','让混合训练自动留下可复盘的记录'),
 ('p','面向跑步与跳绳交替训练、课间运动和日常健身用户，腕动教练把“选择场景—开始训练—查看识别与提醒—停止—复盘”收进四个手表页面。运动状态与时间线在端侧计算，无手机、无云端 AI 服务时仍可完成模拟训练闭环。'),
 ('p','作品基于 openvela 快应用图形框架，采用“小芽”运动伙伴与轻量动画引导训练。首页聚焦当前动作，教练页展示候选与建议，时间线保留片段，同步复盘页承接历史与可选扩展。'),
 ('images',['core_01_home.png','core_02_coach.png']),
 ('p','上图来自本轮 openvela Watch Emulator 的实际 Release 运行。首页展示运行中的跑步状态；界面中的置信度是当前分类输出，并非真人测试准确率。'),
 ('p','价值主张：减少频繁选择运动模式；用连续片段描述一次混合训练；把提醒与复盘放在用户抬腕可见的位置。')]),
('技术实现与 openvela 能力', [
 ('h','从六轴数据到稳定训练片段'),
 ('p','应用内场景 Mock → 16 Hz 六轴样本 → 3 秒滑窗 → IMU 特征 → tiny_classifier → TRL 时间后处理 → 训练片段、风险提示与本地摘要。'),
 ('p','分类标签包括无活动、羽毛球、跳绳、飞鸟、跑步、乒乓球。当前模型为轻量特征/规则分类器；原研究的 CNN-BiLSTM 多尺度模型仅保留迁移接口与资产清单。'),
 ('p','TRL 使用 7 点居中均值、5 点居中中值和 Viterbi 解码，并在固定 32 窗口尾部回溯，冻结稳定前缀。协作式推理与有界队列减少长会话对操作响应的影响；停止训练会取消后续计算。'),
 ('h','系统接口与使用边界'),
 ('p','图形：.ux 页面与系统路由，四页渐进展示；存储：@system.storage 保存摘要、片段与风险事件；振动：@system.vibrator 发出短提醒；健康：@service.health 封装心率、血氧、压力，失败时明确降级。'),
 ('p','传感器：@system.sensor 接入真实加速度与步数诊断。当前 JS 接口缺少 GYRO，完整真机训练会在启动前阻断，不以全零数据继续识别。跨设备：@system.interconnect 发送摘要；可选 @system.velaclaw 总结不可用时退回本地模板。'),
 ('h','两条分别验证的数据链路'),
 ('p','官方模拟器 gRPC 注入与回读验证 ACC、GYRO、心率；应用内 Mock 验证算法和交互。两者共享场景生成逻辑，但本次材料不将“gRPC 注入通过”表述为“应用已通过 JS 消费全部官方六轴数据”。步数采用应用内 Mock。'),
 ('h','隐私设计'),
 ('p','默认端侧计算，本地只保存摘要级记录，不保存原始连续波形。手机导出需主动触发；可选 AI 总结能力的范围与降级路径在 docs/privacy_statement.md 中说明。')]),
('验证结果与评委复现', [
 ('p','本轮验收：Node.js v22.23.1、aiot-toolkit 2.0.5；openvela Vela5 Watch Emulator，原生显示 432×514，采集图适配为 403×480。官方工程接入目标分支为 dev-ai-contest-2026；未宣称本地预编译模拟器镜像由该分支源码构建。'),
 ('p','核心回归 53 项通过，覆盖单位、采样网格、缺失陀螺仪门禁、异步生命周期、停止/诊断互斥、存储和同步状态。运动场景回归全部通过，覆盖六类、混合训练和疲劳提示；这些是合成数据逻辑测试。'),
 ('p','生产 RPK 构建成功。部署后核对设备页面 bundle 与本地 Release SHA-256 一致；动态日志与画面均确认跑步，再执行停止与翻页。四页截图互不重复，严格画面检查通过。官方 Mock 的 ACC、GYRO、心率注入及回读通过。'),
 ('images',['core_03_timeline.png','core_04_sync_review.png']),
 ('p','复现：进入 quickapp/velamotion_coach/，执行 npm ci、npm run test:core、npm run test:motion；直接安装 dist/ 中的签名 Release 即可演示。自行重新发布可在 AIoT-IDE 生成自己的签名，原作者私钥不随仓库分发。'),
 ('p','运行：安装后从首页点击开始，等待 3 秒以上预热；教练页可选跑步/跳绳/混合训练等场景。停止后查看时间线、本地复盘与历史。模拟器在线时运行 npm run demo:capture 可重新采集四页。')]),
('交付清单、AI 协作与后续计划', [
 ('h','评审入口'),
 ('p','源码仓库：https://github.com/rudykon/VelaMotion_Coach\n源码与 RPK：quickapp/velamotion_coach/\n介绍文件：docs/作品介绍.pdf、docs/作品介绍.docx\n视频：artifacts/final_demo/auto_carousel/velamotion_core_demo.mp4\nAI Coding 日志：本次不公开，仅本地保留\n复用 Skill：skills/openvela-watch-acceptance/SKILL.md'),
 ('p','视频是四张实际模拟器画面的短片串联，不是连续交互录屏。完整点击和动态识别过程见验收记录及可复现脚本。源码采用 Apache-2.0；签名私钥、依赖缓存、原始训练数据和论文材料均不纳入作品。'),
 ('h','AI Coding 的具体作用'),
 ('p','AI 协作用于采样边界排查、TRL 增量优化、交互状态检查和模拟器验收，并沉淀可复用验收 Skill。按作者要求，本次不公开历史开发日志。该项官方材料尚未提交，作品运行和打包检查通过不等同于全部参赛要求已满足。'),
 ('h','当前限制与下一步'),
 ('p','当前验证集中于模拟器：真实用户识别准确率、真实手表功耗、振动触感和手机配对端到端效果尚未验证。提醒为规则提示，不具备临床验证。下一步补齐目标设备六轴采样和原生模型桥，在授权真机数据上比较准确率、响应时延和能耗，再验证手机复盘。'),
 ('p','潜在应用包括校园体测训练、个人混合健身记录和腕上轻量教练。当前为参赛原型，尚无营收或真实用户规模数据；推广优先验证“自动记录是否减少操作负担”。'),
 ('h','规则依据（2026-09-20 核对）'),
 ('p','open-vela/docs 的 dev-ai-contest-2026 分支：contest_overview.md、code_submission_guide.md、quickapp/watch_app_track_guide.md、quickapp/quickapp_manual.md、ai_coding_log_guide.md。完整链接与逐项对应见 docs/contest_requirements.md。')])]

pdf=[]
doc=Document()
sec=doc.sections[0]; sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65)
normal=doc.styles['Normal'];normal.font.name='Noto Sans CJK SC';normal.font.size=Pt(10)
normal.element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
normal.paragraph_format.space_after=Pt(7)
for style in ('Title','Heading 1','Heading 2'):
 doc.styles[style].font.name='Noto Sans CJK SC';doc.styles[style].element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC');doc.styles[style].font.color.rgb=RGBColor.from_string('16755A')
for pi,(title,parts) in enumerate(pages):
 if pi: pdf.append(PageBreak());doc.add_page_break()
 pdf.append(Paragraph(escape(title),styles['CNTitle']));doc.add_heading(title,0)
 for kind,value in parts:
  if kind=='images':
   cells=[];table=doc.add_table(rows=1,cols=len(value))
   for j,name in enumerate(value):
    path=ROOT/'artifacts/final_demo/auto_carousel'/name
    cells.append(Image(str(path),width=162,height=193))
    table.cell(0,j).paragraphs[0].add_run().add_picture(str(path),width=Inches(2.15))
   t=Table([cells],colWidths=[242]*len(cells));t.setStyle(TableStyle([('ALIGN',(0,0),(-1,-1),'CENTER'),('BOTTOMPADDING',(0,0),(-1,-1),10)]));pdf.append(t)
  else:
   pdf.append(Paragraph(escape(value).replace('\n','<br/>'),styles['CNHeading' if kind=='h' else 'CN']))
   if kind=='h':doc.add_heading(value,2)
   else:doc.add_paragraph(value)

def footer(c,d):
 c.setStrokeColor(colors.HexColor('#CDE6DC'));c.line(42,37,A4[0]-42,37)
 c.setFont('Chinese',8);c.setFillColor(colors.HexColor('#57766B'))
 c.drawString(42,24,'VelaMotion Coach | DDLqudong · 289 | 2026-09-20')
 c.drawRightString(A4[0]-42,24,str(d.page))
SimpleDocTemplate(str(OUT/'作品介绍.pdf'),pagesize=A4,leftMargin=42,rightMargin=42,topMargin=38,bottomMargin=47,title='腕动教练 - VelaMotion Coach',author='DDLqudong').build(pdf,onFirstPage=footer,onLaterPages=footer)
doc.core_properties.title='腕动教练 - VelaMotion Coach';doc.core_properties.author='DDLqudong'
doc.save(OUT/'作品介绍.docx')
print('Generated docs/作品介绍.pdf and docs/作品介绍.docx')
